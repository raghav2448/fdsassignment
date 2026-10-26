# %% [markdown]
# # Fundamentals of Data Science – Combined Assignment
#
# **Part A – Bank Marketing:** comparative analysis of classification algorithms (Logistic Regression, Decision Tree, KNN, Random Forest).
#
# **Part B – Online Retail:** transaction analysis and comparison of clustering algorithms (K-means, Hierarchical, DBSCAN, Spectral) with PCA.
#
# **How to run:** put the datasets in a folder called `data/` next to this file (or set `BANK_PATH` / `RETAIL_PATH` below), then run all cells.
# Every figure, table and the final Word report (`outputs/Combined_Report.docx`) are produced from the results of this code.

# %%
import os
import re
import csv
import glob
import math
import warnings
from collections import Counter

import numpy as np
import pandas as pd
import matplotlib

try:  # running inside Jupyter?
    get_ipython()
    IN_NOTEBOOK = True
except NameError:
    IN_NOTEBOOK = False
    matplotlib.use("Agg")          # no pop-up windows when run as a plain script
import matplotlib.pyplot as plt

from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler, OneHotEncoder, MultiLabelBinarizer
from sklearn.compose import ColumnTransformer
from sklearn.linear_model import LogisticRegression
from sklearn.tree import DecisionTreeClassifier
from sklearn.neighbors import KNeighborsClassifier, NearestNeighbors
from sklearn.ensemble import RandomForestClassifier
from sklearn.metrics import (accuracy_score, precision_score, recall_score, f1_score,
                             confusion_matrix, ConfusionMatrixDisplay, roc_curve,
                             roc_auc_score, precision_recall_curve,
                             silhouette_score, davies_bouldin_score)
from sklearn.decomposition import PCA
from sklearn.cluster import KMeans, AgglomerativeClustering, DBSCAN, SpectralClustering
from scipy.cluster.hierarchy import linkage, dendrogram

warnings.filterwarnings("ignore")
plt.rcParams.update({"figure.dpi": 100, "axes.grid": True, "grid.alpha": 0.3})

# %% [markdown]
# ## 0. Configuration and helper functions

# %%
# ---------- EDIT THESE ----------
STUDENT_NAME = "Your Name"
STUDENT_ID = "Your Roll No / ID"
COURSE = "Fundamentals of Data Science"

BANK_PATH = None      # e.g. "data/bank-additional-full.csv"  (None = auto-detect inside data/)
RETAIL_PATH = None    # e.g. "data/Market_Basket_Optimisation.csv" (None = auto-detect inside data/)
# --------------------------------

SEED = 42
DATA_DIR = "data"
OUT_DIR = "outputs"
FIG_DIR = os.path.join(OUT_DIR, "figures")
TAB_DIR = os.path.join(OUT_DIR, "tables")
for d in (FIG_DIR, TAB_DIR):
    os.makedirs(d, exist_ok=True)

FIGS = {}      # name -> image path (used by the report builder)
R = {}         # all numbers/tables needed by the report are stored here


def find_file(keywords, override=None):
    """Locate a data file whose name contains one of the keywords."""
    override = override or None
    if override and os.path.exists(override):
        return override
    for folder in (DATA_DIR, "."):
        for p in sorted(glob.glob(os.path.join(folder, "*"))):
            if p.lower().endswith((".csv", ".xlsx", ".xls")) and any(k in os.path.basename(p).lower() for k in keywords):
                return p
    return None


def show(obj, title=None):
    """Display a DataFrame/Series nicely in Jupyter, print in a script."""
    if title:
        print(f"\n=== {title} ===")
    if IN_NOTEBOOK:
        from IPython.display import display
        display(obj)
    else:
        print(obj.to_string() if isinstance(obj, (pd.DataFrame, pd.Series)) else obj)


def save_fig(name):
    path = os.path.join(FIG_DIR, name + ".png")
    plt.tight_layout()
    plt.savefig(path, dpi=150, bbox_inches="tight")
    if IN_NOTEBOOK:
        plt.show()
    else:
        plt.close()
    FIGS[name] = path


def save_table(df, name):
    df.to_csv(os.path.join(TAB_DIR, name + ".csv"))


BANK_PATH = find_file(["bank"], os.environ.get("BANK_CSV", BANK_PATH))
RETAIL_PATH = find_file(["market", "basket", "retail", "transaction", "groceries"], os.environ.get("RETAIL_CSV", RETAIL_PATH))
RUN_BANK, RUN_RETAIL = BANK_PATH is not None, RETAIL_PATH is not None
print("Bank file  :", BANK_PATH)
print("Retail file:", RETAIL_PATH)
if not (RUN_BANK or RUN_RETAIL):
    raise FileNotFoundError("No dataset found. Put the CSV files in a 'data/' folder or set BANK_PATH / RETAIL_PATH.")

# %% [markdown]
# # PART A – Bank Marketing: Classification
# ## Task 1 – Data understanding and preprocessing

# %%
if RUN_BANK:
    bank_raw = pd.read_csv(BANK_PATH, sep=";")                     # semicolon-delimited
    bank_raw.columns = bank_raw.columns.str.strip().str.replace('"', "", regex=False)

    show(bank_raw.head(10), "First 10 records")
    print("\nDataset dimensions (rows, columns):", bank_raw.shape)
    show(bank_raw.dtypes.rename("dtype").to_frame(), "Data types")
    show(bank_raw.describe().T, "Descriptive statistics – numerical attributes")
    show(bank_raw.describe(include="object").T, "Descriptive statistics – categorical attributes")
    R["bank_shape_raw"] = bank_raw.shape

# %%
if RUN_BANK:
    TARGET = "y"
    num_cols_all = bank_raw.select_dtypes(include="number").columns.tolist()
    cat_cols_all = [c for c in bank_raw.select_dtypes(exclude="number").columns if c != TARGET]
    print("Target attribute      :", TARGET)
    print("Numerical attributes  :", num_cols_all)
    print("Categorical attributes:", cat_cols_all)

    # Missing, 'unknown' and duplicate records
    missing = bank_raw.isna().sum()
    unknown = pd.Series({c: int((bank_raw[c].astype(str).str.lower() == "unknown").sum()) for c in cat_cols_all})
    n_dups = int(bank_raw.duplicated().sum())
    show(missing[missing > 0] if (missing > 0).any() else "No true missing (NaN) values.", "Missing values")
    show(unknown[unknown > 0].rename("'unknown' entries").to_frame(), "'unknown' entries per categorical column")
    print("Duplicate records:", n_dups)

    bank = bank_raw.drop_duplicates().reset_index(drop=True)      # remove exact duplicates
    bank["subscribed"] = (bank[TARGET].str.lower() == "yes").astype(int)
    tdist = bank[TARGET].value_counts()
    show(tdist.to_frame("count").assign(percent=(tdist / len(bank) * 100).round(2)), "Target distribution")
    R.update(missing_total=int(missing.sum()), unknown=unknown[unknown > 0].to_dict(), n_dups=n_dups,
             bank_shape_clean=bank.shape, num_cols_all=num_cols_all, cat_cols_all=cat_cols_all)

# %% [markdown]
# **Preprocessing decisions**
# * `unknown` is kept as its own category (it can itself be informative, e.g. an unknown default status).
# * Exact duplicate rows are removed.
# * **`duration` is excluded** from all models: it is only known after the call ends, so using it would leak the outcome.
# * Categorical variables → one-hot encoding; numerical variables → standardisation (needed by KNN and Logistic Regression).
# * 80:20 stratified split. The encoder/scaler are fitted **on the training set only** to avoid data leakage.

# %%
if RUN_BANK:
    DROP = [c for c in ["duration"] if c in bank.columns]
    num_feats = [c for c in num_cols_all if c not in DROP]
    cat_feats = cat_cols_all
    X = bank[num_feats + cat_feats]
    y = bank["subscribed"]

    X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.20, stratify=y, random_state=SEED)

    preprocessor = ColumnTransformer([
        ("num", StandardScaler(), num_feats),
        ("cat", OneHotEncoder(handle_unknown="ignore", sparse_output=False), cat_feats)])
    Xtr = preprocessor.fit_transform(X_train)
    Xte = preprocessor.transform(X_test)
    feature_names = preprocessor.get_feature_names_out()

    print("Excluded from modelling:", DROP)
    print("Predictors after encoding:", Xtr.shape[1])
    print(f"Train: {Xtr.shape[0]} rows | Test: {Xte.shape[0]} rows")
    print("Subscription rate  train: %.2f%%  test: %.2f%%" % (y_train.mean() * 100, y_test.mean() * 100))
    R.update(n_train=len(y_train), n_test=len(y_test), n_features_encoded=Xtr.shape[1], dropped=DROP,
             num_feats=num_feats, cat_feats=cat_feats)

# %% [markdown]
# ## Task 2 – Exploratory data analysis

# %%
if RUN_BANK:
    n_yes = int(bank["subscribed"].sum()); n_no = len(bank) - n_yes
    eda_counts = pd.DataFrame({"Customers": [n_yes, n_no],
                               "Percentage": [round(n_yes / len(bank) * 100, 2), round(n_no / len(bank) * 100, 2)]},
                              index=["Subscribed (yes)", "Not subscribed (no)"])
    show(eda_counts, "Subscribed vs not subscribed")
    R.update(n_yes=n_yes, n_no=n_no, pct_yes=n_yes / len(bank) * 100, eda_counts=eda_counts)

    fig, ax = plt.subplots(figsize=(6, 4.5))
    bars = ax.bar(["No", "Yes"], [n_no, n_yes], color=["#9aa5b1", "#2a7de1"])
    for b, v in zip(bars, [n_no, n_yes]):
        ax.text(b.get_x() + b.get_width() / 2, v, f"{v:,}\n({v / len(bank) * 100:.1f}%)", ha="center", va="bottom")
    ax.set_title("Distribution of the Target Variable (Term Deposit Subscription)")
    ax.set_xlabel("Subscribed to term deposit?"); ax.set_ylabel("Number of customers")
    ax.set_ylim(0, max(n_no, n_yes) * 1.15)
    save_fig("A1_target_distribution")

# %%
if RUN_BANK:
    MONTHS = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
    rate_cols = [c for c in ["job", "marital", "education", "contact", "month", "poutcome"] if c in bank.columns]
    overall = bank["subscribed"].mean() * 100

    fig, axes = plt.subplots(2, 3, figsize=(17, 9))
    for ax, col in zip(axes.ravel(), rate_cols):
        r = bank.groupby(col)["subscribed"].mean().mul(100)
        if col == "month":
            r = r.reindex([m for m in MONTHS if m in r.index])
            r.plot.bar(ax=ax, color="#2a7de1"); ax.axhline(overall, color="red", ls="--", lw=1)
            ax.set_ylabel("Subscription rate (%)"); ax.set_xlabel(col)
        else:
            r.sort_values().plot.barh(ax=ax, color="#2a7de1"); ax.axvline(overall, color="red", ls="--", lw=1)
            ax.set_xlabel("Subscription rate (%)"); ax.set_ylabel(col)
        ax.set_title(f"Subscription rate by {col}")
    for ax in axes.ravel()[len(rate_cols):]:
        ax.axis("off")
    fig.suptitle(f"Subscription rate across customer / campaign attributes (red dashed line = overall {overall:.1f}%)", y=1.01)
    save_fig("A2_subscription_rate_by_category")

    # category-level table used for observations
    rows = []
    for col in rate_cols:
        g = bank.groupby(col)["subscribed"].agg(["size", "mean"])
        for cat, rr in g.iterrows():
            rows.append((col, cat, int(rr["size"]), rr["mean"] * 100))
    cat_rates = pd.DataFrame(rows, columns=["attribute", "category", "n", "rate_pct"])
    cat_rates["lift"] = cat_rates["rate_pct"] / overall
    R["cat_rates"] = cat_rates

# %%
if RUN_BANK:
    cand_num = [c for c in ["age", "campaign", "previous", "euribor3m", "balance", "cons.conf.idx", "emp.var.rate"] if c in bank.columns]
    plot_num = cand_num[:4]
    fig, axes = plt.subplots(1, len(plot_num), figsize=(4.2 * len(plot_num), 4.5))
    axes = np.atleast_1d(axes)
    for ax, col in zip(axes, plot_num):
        data = [bank.loc[bank["subscribed"] == 0, col], bank.loc[bank["subscribed"] == 1, col]]
        ax.boxplot(data, labels=["No", "Yes"], showfliers=False, patch_artist=True,
                   boxprops=dict(facecolor="#cfe0f7"), medianprops=dict(color="red"))
        ax.set_title(f"{col} vs subscription"); ax.set_xlabel("Subscribed?"); ax.set_ylabel(col)
    fig.suptitle("Numerical attributes by subscription outcome (outliers hidden)", y=1.02)
    save_fig("A3_numeric_vs_outcome")

    num_summary = bank.groupby("subscribed")[[c for c in num_feats]].mean().T.rename(columns={0: "mean (No)", 1: "mean (Yes)"})
    num_summary["std (all)"] = bank[num_feats].std()
    num_summary["std. difference"] = (num_summary["mean (Yes)"] - num_summary["mean (No)"]) / num_summary["std (all)"]
    show(num_summary.round(3), "Mean of numerical attributes by outcome")
    R["num_summary"] = num_summary

    if "age" in bank.columns:
        bank["age_group"] = pd.cut(bank["age"], [0, 25, 35, 45, 55, 65, 120], labels=["<=25", "26-35", "36-45", "46-55", "56-65", "65+"])
        ar = bank.groupby("age_group", observed=True)["subscribed"].mean().mul(100)
        fig, ax = plt.subplots(figsize=(7, 4.5))
        ar.plot.bar(ax=ax, color="#2a7de1"); ax.axhline(overall, color="red", ls="--")
        ax.set_title("Subscription rate by age group"); ax.set_xlabel("Age group"); ax.set_ylabel("Subscription rate (%)")
        save_fig("A4_subscription_by_age_group")
        R["age_rates"] = ar

# %%
if RUN_BANK:
    # --- three (plus one) data-driven observations ---
    obs = []
    obs.append(f"The target is imbalanced: only {R['n_yes']:,} of {len(bank):,} customers ({R['pct_yes']:.1f}%) subscribed, "
               f"so a model that always predicts 'no' would already reach about {100 - R['pct_yes']:.1f}% accuracy – accuracy alone is misleading.")
    big = cat_rates[cat_rates["n"] >= 100].sort_values("rate_pct", ascending=False)
    t1 = big.iloc[0]
    obs.append(f"Customers with {t1['attribute']} = '{t1['category']}' subscribe at {t1['rate_pct']:.1f}% "
               f"({t1['lift']:.1f}x the overall rate, n={t1['n']:,}).")
    t2 = big[big["attribute"] != t1["attribute"]].iloc[0]
    obs.append(f"A second high-yield segment is {t2['attribute']} = '{t2['category']}' with {t2['rate_pct']:.1f}% "
               f"({t2['lift']:.1f}x the overall rate, n={t2['n']:,}); the weakest large segment is "
               f"{big.iloc[::-1].iloc[0]['attribute']} = '{big.iloc[::-1].iloc[0]['category']}' at {big.iloc[::-1].iloc[0]['rate_pct']:.1f}%.")
    top_num = num_summary["std. difference"].abs().sort_values(ascending=False).index[0]
    d = num_summary.loc[top_num]
    direction = "higher" if d["mean (Yes)"] > d["mean (No)"] else "lower"
    obs.append(f"Among numerical attributes, '{top_num}' separates the classes most: subscribers have a {direction} average "
               f"({d['mean (Yes)']:.2f} vs {d['mean (No)']:.2f} for non-subscribers).")
    R["observations"] = obs
    for i, o in enumerate(obs, 1):
        print(f"Observation {i}: {o}\n")

# %% [markdown]
# ## Task 3 – Classification model development
# Same train/test split, same preprocessing and same target definition for all four models.
# Hyper-parameters are fixed (no tuning on the test set): Logistic Regression (`max_iter=1000`), Decision Tree (`max_depth=6, min_samples_leaf=50`),
# KNN (`k=15`), Random Forest (`200 trees, min_samples_leaf=5`).

# %%
def make_models(balanced=False):
    cw = "balanced" if balanced else None
    return {
        "Logistic Regression": LogisticRegression(max_iter=1000, class_weight=cw, random_state=SEED),
        "Decision Tree": DecisionTreeClassifier(max_depth=6, min_samples_leaf=50, class_weight=cw, random_state=SEED),
        "KNN": KNeighborsClassifier(n_neighbors=15),
        "Random Forest": RandomForestClassifier(n_estimators=200, min_samples_leaf=5, class_weight=cw, n_jobs=-1, random_state=SEED),
    }


def evaluate_models(models, Xtr, ytr, Xte, yte):
    rows, store = [], {}
    for name, model in models.items():
        model.fit(Xtr, ytr)                                  # train
        pred = model.predict(Xte)                            # predict
        proba = model.predict_proba(Xte)[:, 1]
        cm = confusion_matrix(yte, pred)
        rows.append({"Model": name,
                     "Accuracy": accuracy_score(yte, pred),
                     "Precision": precision_score(yte, pred, zero_division=0),
                     "Recall": recall_score(yte, pred, zero_division=0),
                     "F1-score": f1_score(yte, pred, zero_division=0),
                     "ROC-AUC": roc_auc_score(yte, proba)})
        store[name] = {"model": model, "pred": pred, "proba": proba, "cm": cm}
        tn, fp, fn, tp = cm.ravel()
        print(f"{name:20s} acc={rows[-1]['Accuracy']:.4f} prec={rows[-1]['Precision']:.4f} "
              f"rec={rows[-1]['Recall']:.4f} f1={rows[-1]['F1-score']:.4f} auc={rows[-1]['ROC-AUC']:.4f} | TN={tn} FP={fp} FN={fn} TP={tp}")
    return pd.DataFrame(rows).set_index("Model"), store


if RUN_BANK:
    results, store = evaluate_models(make_models(False), Xtr, y_train, Xte, y_test)
    save_table(results.round(4), "A_model_comparison")
    R["results"] = results

# %% [markdown]
# ## Task 4 – Model comparison and visualisation

# %%
if RUN_BANK:
    show(results.round(4), "Final performance comparison table (test set, default 0.5 threshold)")

    ax = results[["Accuracy", "Precision", "Recall", "F1-score"]].plot.bar(figsize=(10, 5.5), width=0.8, edgecolor="black")
    for c in ax.containers:
        ax.bar_label(c, fmt="%.2f", fontsize=8, padding=2)
    ax.set_title("Comparison of Classification Models – Accuracy, Precision, Recall and F1-score")
    ax.set_xlabel("Model"); ax.set_ylabel("Score"); ax.set_ylim(0, 1.08)
    plt.xticks(rotation=0); ax.legend(loc="upper right", ncol=4, fontsize=9)
    save_fig("A5_model_comparison_grouped_bar")

# %%
if RUN_BANK:
    fig, axes = plt.subplots(2, 2, figsize=(11, 9))
    for ax, (name, s) in zip(axes.ravel(), store.items()):
        ConfusionMatrixDisplay.from_predictions(y_test, s["pred"], ax=ax, cmap="Blues", display_labels=["No", "Yes"],
                                                values_format="d", colorbar=False)
        ax.set_title(f"Confusion matrix – {name}"); ax.set_xlabel("Predicted label"); ax.set_ylabel("True label"); ax.grid(False)
    save_fig("A6_confusion_matrices")

    cm_rows = []
    for name, s in store.items():
        tn, fp, fn, tp = s["cm"].ravel()
        cm_rows.append({"Model": name, "TN": tn, "FP": fp, "FN": fn, "TP": tp,
                        "Actual subscribers": tp + fn, "Subscribers found (recall %)": round(tp / (tp + fn) * 100, 1),
                        "Subscribers missed": fn, "Wasted calls (FP)": fp})
    cm_table = pd.DataFrame(cm_rows).set_index("Model")
    show(cm_table, "Classification errors and ability to find actual subscribers")
    save_table(cm_table, "A_confusion_summary")
    R["cm_table"] = cm_table

# %%
if RUN_BANK:
    fig, ax = plt.subplots(figsize=(7.5, 6.5))
    for name, s in store.items():
        fpr, tpr, _ = roc_curve(y_test, s["proba"])
        ax.plot(fpr, tpr, lw=2, label=f"{name} (AUC = {results.loc[name, 'ROC-AUC']:.3f})")
    ax.plot([0, 1], [0, 1], "k--", label="Random guess (AUC = 0.500)")
    ax.set_title("ROC Curves of the Classifiers"); ax.set_xlabel("False Positive Rate"); ax.set_ylabel("True Positive Rate (Recall)")
    ax.legend(loc="lower right")
    save_fig("A7_roc_curves")

    # cumulative gains: how many subscribers are found if we call the best-scored x% of customers
    fig, ax = plt.subplots(figsize=(7.5, 6.5))
    yt = y_test.values
    for name, s in store.items():
        order = np.argsort(-s["proba"])
        gains = np.cumsum(yt[order]) / yt.sum()
        ax.plot(np.arange(1, len(yt) + 1) / len(yt) * 100, gains * 100, lw=2, label=name)
    ax.plot([0, 100], [0, 100], "k--", label="Random calling")
    ax.set_title("Cumulative Gains – subscribers captured vs customers contacted")
    ax.set_xlabel("Customers contacted, ranked by model score (%)"); ax.set_ylabel("Actual subscribers captured (%)")
    ax.legend(loc="lower right")
    save_fig("A8_cumulative_gains")

# %% [markdown]
# ### Choosing the best model
# Accuracy is **not** used for selection. Because a missed subscriber (false negative) is a lost deposit while a false positive is only one extra phone call,
# the selection score is the mean of **ROC-AUC, F1-score and Recall**:
# `score = (ROC-AUC + F1 + Recall) / 3`.

# %%
if RUN_BANK:
    results["Selection score"] = (results["ROC-AUC"] + results["F1-score"] + results["Recall"]) / 3
    best_model = results["Selection score"].idxmax()
    show(results.sort_values("Selection score", ascending=False).round(4), "Ranking by selection score")
    print("\nBEST MODEL:", best_model)

    best_proba = store[best_model]["proba"]
    # (a) top 20% of customers
    n_top = math.ceil(0.20 * len(yt)); order = np.argsort(-best_proba)[:n_top]
    capture20 = yt[order].sum() / yt.sum()
    # (b) threshold that captures >= 60% of subscribers
    prec_c, rec_c, thr_c = precision_recall_curve(yt, best_proba)
    idx = np.where(rec_c[:-1] >= 0.60)[0]
    i = idx[-1] if len(idx) else 0
    thr60, prec60, rec60, frac60 = thr_c[i], prec_c[i], rec_c[i], (best_proba >= thr_c[i]).mean()
    R.update(best_model=best_model, capture20=capture20, lift20=capture20 / 0.20, thr60=thr60, prec60=prec60,
             rec60=rec60, frac60=frac60, results_scored=results.copy())
    print(f"Calling the top 20% of customers (by {best_model} score) captures {capture20 * 100:.1f}% of all subscribers "
          f"(lift {capture20 / 0.2:.2f}x).")
    print(f"Threshold {thr60:.3f} gives recall {rec60 * 100:.1f}% with precision {prec60 * 100:.1f}% "
          f"while contacting {frac60 * 100:.1f}% of customers.")

# %% [markdown]
# ### Extension – effect of class imbalance handling
# Logistic Regression, Decision Tree and Random Forest are re-trained with `class_weight="balanced"` (KNN has no such option and is unchanged).
# This is **not** used for the primary selection; it shows the precision/recall trade-off the bank can choose from.

# %%
if RUN_BANK:
    results_bal, store_bal = evaluate_models(make_models(True), Xtr, y_train, Xte, y_test)
    show(results_bal.round(4), "Extension: class_weight='balanced'")
    save_table(results_bal.round(4), "A_model_comparison_balanced")
    R["results_bal"] = results_bal

# %% [markdown]
# # PART B – Online Retail: Transaction Clustering
# ## Objectives 1–4: data understanding, cleaning, descriptive statistics and visualisation

# %%
NOISE_ITEMS = {"postage", "dotcom postage", "manual", "carriage", "bank charges", "amazon fee", "discount", "samples", "cruk commission"}


def _clean_item(x):
    x = re.sub(r"\s+", " ", str(x)).strip().lower()
    return None if x in ("", "nan", "none") else x


def load_transactions(path):
    """Return (list_of_item_lists, info_dict). Handles basket-per-row files and invoice/item (long) files."""
    ext = os.path.splitext(path)[1].lower()
    info = {"file": os.path.basename(path)}
    if ext in (".xlsx", ".xls"):
        raw = pd.read_excel(path); fmt = "long"
    else:
        with open(path, encoding="utf-8", errors="ignore") as f:
            first = f.readline().lower()
        fmt = "long" if any(k in first for k in ("invoice", "description", "stockcode", "transaction", "customer")) else "basket"
        if fmt == "long":
            try:
                raw = pd.read_csv(path, encoding="utf-8")
            except UnicodeDecodeError:
                raw = pd.read_csv(path, encoding="latin1")
    info["format"] = fmt

    if fmt == "basket":                                    # every row = one transaction, items spread over columns
        with open(path, encoding="utf-8", errors="ignore", newline="") as f:
            rows = list(csv.reader(f))
        info["raw_rows"] = len(rows)
        tx = [sorted({i for i in (_clean_item(v) for v in row) if i}) for row in rows]
    else:                                                  # one row = one purchased line
        cols = {c.lower(): c for c in raw.columns}
        inv = next((cols[c] for c in cols if any(k in c for k in ("invoice", "transaction", "order", "basket", "bill"))), None)
        itm = next((cols[c] for c in cols if any(k in c for k in ("description", "item", "product"))), None)
        if inv is None or itm is None:
            raise ValueError(f"Could not find invoice/item columns in {list(raw.columns)}")
        info["raw_rows"] = len(raw)
        raw = raw.dropna(subset=[inv, itm])
        raw = raw[~raw[inv].astype(str).str.upper().str.startswith("C")]          # cancellations
        qty = next((cols[c] for c in cols if "quantity" in c), None)
        if qty is not None:
            raw = raw[raw[qty] > 0]
        raw[itm] = raw[itm].map(_clean_item)
        raw = raw.dropna(subset=[itm]); raw = raw[~raw[itm].isin(NOISE_ITEMS)]
        tx = [sorted(set(v)) for v in raw.groupby(inv)[itm].apply(list)]
    tx = [t for t in tx if len(t) > 0]
    info["n_transactions"] = len(tx)
    return tx, info


if RUN_RETAIL:
    transactions, retail_info = load_transactions(RETAIL_PATH)
    print(retail_info)
    print("Example transactions:")
    for t in transactions[:5]:
        print("  ", t)
    R["retail_info"] = retail_info

# %%
if RUN_RETAIL:
    basket_sizes = pd.Series([len(t) for t in transactions], name="basket size")
    item_counts = Counter(i for t in transactions for i in t)          # number of transactions containing the item
    n_tx = len(transactions)
    stats = basket_sizes.describe().round(2).to_frame("value")
    show(stats, "Transaction-size statistics (items per transaction)")
    print(f"Transactions: {n_tx:,} | Unique products: {len(item_counts):,} | Most common basket size: {int(basket_sizes.mode()[0])}")

    top_items = pd.DataFrame(item_counts.most_common(15), columns=["Product", "Transactions"])
    top_items["Support (%)"] = (top_items["Transactions"] / n_tx * 100).round(2)
    show(top_items, "Top 15 most frequent products")
    save_table(top_items, "B_top_products")
    R.update(n_tx=n_tx, n_items=len(item_counts), basket_stats=basket_sizes.describe(), top_items=top_items,
             mode_size=int(basket_sizes.mode()[0]), single_item_share=float((basket_sizes == 1).mean() * 100))

    fig, ax = plt.subplots(figsize=(9, 5.5))
    ax.barh(top_items["Product"][::-1], top_items["Transactions"][::-1], color="#2a7de1")
    ax.set_title("Top 15 Most Frequently Purchased Products"); ax.set_xlabel("Number of transactions containing the product"); ax.set_ylabel("Product")
    save_fig("B1_top_products")

    fig, ax = plt.subplots(figsize=(8, 5))
    ax.hist(basket_sizes, bins=range(1, int(basket_sizes.max()) + 2), color="#2a7de1", edgecolor="white", align="left")
    ax.axvline(basket_sizes.mean(), color="red", ls="--", label=f"Mean = {basket_sizes.mean():.2f}")
    ax.axvline(basket_sizes.median(), color="green", ls=":", label=f"Median = {basket_sizes.median():.0f}")
    ax.set_title("Distribution of Transaction Size"); ax.set_xlabel("Number of items in the transaction"); ax.set_ylabel("Number of transactions"); ax.legend()
    save_fig("B2_transaction_size_distribution")

# %% [markdown]
# ## Objectives 5–6: feature vectors and PCA
# Each transaction becomes a **binary vector** (1 = product present). To keep the matrix manageable only the `TOP_N_ITEMS` most frequent products are used.
# Because hierarchical and spectral clustering need an n×n matrix, a random sample of at most `MAX_SAMPLES` transactions is clustered.
# Clustering is performed on the first `N_PCA` principal components; the first two are used for plotting.

# %%
TOP_N_ITEMS = 100
MAX_SAMPLES = 3000
N_PCA = 10

if RUN_RETAIL:
    keep = [i for i, _ in item_counts.most_common(TOP_N_ITEMS)]
    coverage = sum(item_counts[i] for i in keep) / sum(item_counts.values()) * 100
    mlb = MultiLabelBinarizer(classes=keep)
    Xb = mlb.fit_transform(transactions).astype(float)
    nonempty = Xb.sum(axis=1) > 0
    Xb, tx_kept = Xb[nonempty], [t for t, k in zip(transactions, nonempty) if k]

    rng = np.random.RandomState(SEED)
    if len(Xb) > MAX_SAMPLES:
        sel = np.sort(rng.choice(len(Xb), MAX_SAMPLES, replace=False))
        Xs = Xb[sel]
    else:
        sel = np.arange(len(Xb)); Xs = Xb
    print(f"Feature matrix: {Xb.shape} (top {len(keep)} products cover {coverage:.1f}% of all item occurrences)")
    print(f"Clustering sample: {Xs.shape}")

    pca_full = PCA(random_state=SEED).fit(Xs)
    cum = np.cumsum(pca_full.explained_variance_ratio_) * 100
    pca2 = PCA(n_components=2, random_state=SEED).fit(Xs); P2 = pca2.transform(Xs)
    pcaC = PCA(n_components=N_PCA, random_state=SEED).fit(Xs); Z = pcaC.transform(Xs)
    var2, varC = pca2.explained_variance_ratio_.sum() * 100, pcaC.explained_variance_ratio_.sum() * 100
    print(f"Variance explained: 2 PCs = {var2:.1f}% | {N_PCA} PCs = {varC:.1f}%")
    R.update(feat_shape=Xb.shape, coverage=coverage, sample_shape=Xs.shape, var2=var2, varC=varC, n_pca=N_PCA,
             pc1=pca2.explained_variance_ratio_[0] * 100, pc2=pca2.explained_variance_ratio_[1] * 100,
             n_for_80=int(np.searchsorted(cum, 80) + 1))

    fig, axes = plt.subplots(1, 2, figsize=(13, 5))
    axes[0].plot(range(1, 31), pca_full.explained_variance_ratio_[:30] * 100, "o-")
    axes[0].set_title("PCA Scree Plot"); axes[0].set_xlabel("Principal component"); axes[0].set_ylabel("Variance explained (%)")
    axes[1].plot(range(1, len(cum) + 1), cum, "o-", ms=3); axes[1].axhline(80, color="red", ls="--", label="80%")
    axes[1].axvline(N_PCA, color="green", ls=":", label=f"{N_PCA} PCs used")
    axes[1].set_title("Cumulative Explained Variance"); axes[1].set_xlabel("Number of components"); axes[1].set_ylabel("Cumulative variance (%)"); axes[1].legend()
    save_fig("B3_pca_variance")

    fig, ax = plt.subplots(figsize=(7, 6))
    ax.scatter(P2[:, 0], P2[:, 1], s=8, alpha=0.4, color="#2a7de1")
    ax.set_title("Transactions projected onto the first two principal components")
    ax.set_xlabel(f"PC1 ({pca2.explained_variance_ratio_[0] * 100:.1f}% variance)"); ax.set_ylabel(f"PC2 ({pca2.explained_variance_ratio_[1] * 100:.1f}% variance)")
    save_fig("B4_pca_scatter")

# %% [markdown]
# ## Objective 7–8: clustering algorithms and evaluation
# All four algorithms are evaluated on the same PCA space with the **silhouette coefficient** (higher is better, max 1)
# and the **Davies–Bouldin index** (lower is better, min 0). For DBSCAN, noise points are excluded when computing the metrics.

# %%
def cluster_metrics(Zm, labels):
    mask = labels != -1
    k = len(set(labels[mask]))
    if k < 2 or mask.sum() <= k:
        return np.nan, np.nan, k, int((~mask).sum())
    return (silhouette_score(Zm[mask], labels[mask]), davies_bouldin_score(Zm[mask], labels[mask]), k, int((~mask).sum()))


if RUN_RETAIL:
    # ---- K-means: choose k ----
    km_rows = []
    for k in range(2, 9):
        km = KMeans(n_clusters=k, n_init=10, random_state=SEED).fit(Z)
        s, d, _, _ = cluster_metrics(Z, km.labels_)
        km_rows.append({"k": k, "inertia": km.inertia_, "silhouette": s, "davies_bouldin": d})
    km_tab = pd.DataFrame(km_rows).set_index("k")
    show(km_tab.round(4), "K-means for k = 2..8")
    K_BEST = int(km_tab["silhouette"].idxmax())
    print("Selected k (max silhouette):", K_BEST)

    fig, axes = plt.subplots(1, 3, figsize=(16, 4.5))
    axes[0].plot(km_tab.index, km_tab["inertia"], "o-"); axes[0].set_title("K-means Elbow Plot"); axes[0].set_xlabel("k"); axes[0].set_ylabel("Inertia (WCSS)")
    axes[1].plot(km_tab.index, km_tab["silhouette"], "o-", color="green"); axes[1].set_title("Silhouette vs k"); axes[1].set_xlabel("k"); axes[1].set_ylabel("Silhouette")
    axes[2].plot(km_tab.index, km_tab["davies_bouldin"], "o-", color="red"); axes[2].set_title("Davies–Bouldin vs k"); axes[2].set_xlabel("k"); axes[2].set_ylabel("Davies–Bouldin")
    save_fig("B5_kmeans_selection")
    R.update(km_tab=km_tab, k_best=K_BEST)

    labels_all, params = {}, {}
    labels_all["K-means"] = KMeans(n_clusters=K_BEST, n_init=10, random_state=SEED).fit_predict(Z)
    params["K-means"] = f"k={K_BEST}"

# %%
if RUN_RETAIL:
    # ---- Hierarchical (Ward) ----
    labels_all["Hierarchical"] = AgglomerativeClustering(n_clusters=K_BEST, linkage="ward").fit_predict(Z)
    params["Hierarchical"] = f"Ward linkage, k={K_BEST}"

    sub = Z[np.random.RandomState(SEED).choice(len(Z), min(200, len(Z)), replace=False)]
    fig, ax = plt.subplots(figsize=(11, 5))
    dendrogram(linkage(sub, "ward"), truncate_mode="lastp", p=25, ax=ax, leaf_rotation=90)
    ax.set_title("Hierarchical Clustering Dendrogram (Ward, 200 sampled transactions, last 25 merges)")
    ax.set_xlabel("Cluster (size in brackets)"); ax.set_ylabel("Ward distance")
    save_fig("B6_dendrogram")

    # ---- DBSCAN: grid search over eps / min_samples ----
    best_db, db_rows = None, []
    for ms in (5, 10, 20):
        dist = NearestNeighbors(n_neighbors=ms).fit(Z).kneighbors(Z)[0][:, -1]
        if ms == 10:
            kdist = np.sort(dist)
        for q in (40, 50, 60, 70, 80, 90, 95):
            eps = float(np.percentile(dist, q))
            lab = DBSCAN(eps=eps, min_samples=ms).fit_predict(Z)
            s, d, k, noise = cluster_metrics(Z, lab)
            db_rows.append({"eps": round(eps, 4), "min_samples": ms, "clusters": k, "noise": noise, "silhouette": s, "davies_bouldin": d})
            if k >= 2 and noise / len(Z) <= 0.5 and not np.isnan(s) and (best_db is None or s > best_db[0]):
                best_db = (s, eps, ms, lab)
    db_tab = pd.DataFrame(db_rows)
    show(db_tab.round(4), "DBSCAN parameter search")
    if best_db is not None:
        labels_all["DBSCAN"] = best_db[3]; params["DBSCAN"] = f"eps={best_db[1]:.3f}, min_samples={best_db[2]}"
    else:
        print("DBSCAN did not find >=2 clusters with <=50% noise for any setting – reported as not applicable.")

    fig, ax = plt.subplots(figsize=(7, 4.5))
    ax.plot(kdist); ax.set_title("DBSCAN k-distance plot (min_samples = 10)")
    ax.set_xlabel("Transactions sorted by distance"); ax.set_ylabel("Distance to 10th nearest neighbour")
    save_fig("B7_dbscan_kdistance")
    R["db_tab"] = db_tab

# %%
if RUN_RETAIL:
    # ---- Spectral ----
    try:
        labels_all["Spectral"] = SpectralClustering(n_clusters=K_BEST, affinity="nearest_neighbors", n_neighbors=15,
                                                    assign_labels="kmeans", random_state=SEED).fit_predict(Z)
        params["Spectral"] = f"k={K_BEST}, 15-NN affinity graph"
    except Exception as e:
        print("Spectral clustering failed:", e)

    rows = []
    for name in ["K-means", "Hierarchical", "DBSCAN", "Spectral"]:
        if name in labels_all:
            s, d, k, noise = cluster_metrics(Z, labels_all[name])
            rows.append({"Algorithm": name, "Parameters": params[name], "Clusters": k, "Noise points": noise,
                         "Silhouette (higher=better)": s, "Davies-Bouldin (lower=better)": d})
        else:
            rows.append({"Algorithm": name, "Parameters": "n/a", "Clusters": 0, "Noise points": 0,
                         "Silhouette (higher=better)": np.nan, "Davies-Bouldin (lower=better)": np.nan})
    comp = pd.DataFrame(rows).set_index("Algorithm")
    valid = comp.dropna(subset=["Silhouette (higher=better)"])
    comp["Rank (silhouette)"] = comp["Silhouette (higher=better)"].rank(ascending=False)
    comp["Rank (DB)"] = comp["Davies-Bouldin (lower=better)"].rank(ascending=True)
    comp["Average rank"] = comp[["Rank (silhouette)", "Rank (DB)"]].mean(axis=1)
    show(comp.round(4), "Clustering comparison")
    save_table(comp.round(4), "B_clustering_comparison")
    best_sil = valid["Silhouette (higher=better)"].idxmax(); best_dbi = valid["Davies-Bouldin (lower=better)"].idxmin()
    best_overall = comp["Average rank"].idxmin()
    print(f"Best silhouette: {best_sil} | Best Davies-Bouldin: {best_dbi} | Best average rank: {best_overall}")
    R.update(comp=comp, best_sil=best_sil, best_dbi=best_dbi, best_overall=best_overall, cluster_params=params)

    fig, axes = plt.subplots(1, 2, figsize=(12, 4.5))
    comp["Silhouette (higher=better)"].plot.bar(ax=axes[0], color="#2a7de1", rot=0); axes[0].set_title("Silhouette coefficient (higher is better)"); axes[0].set_ylabel("Silhouette"); axes[0].set_xlabel("Algorithm")
    comp["Davies-Bouldin (lower=better)"].plot.bar(ax=axes[1], color="#e1692a", rot=0); axes[1].set_title("Davies–Bouldin index (lower is better)"); axes[1].set_ylabel("Davies–Bouldin"); axes[1].set_xlabel("Algorithm")
    save_fig("B8_clustering_metrics")

    fig, axes = plt.subplots(2, 2, figsize=(13, 10))
    for ax, name in zip(axes.ravel(), ["K-means", "Hierarchical", "DBSCAN", "Spectral"]):
        if name in labels_all:
            lab = labels_all[name]
            ax.scatter(P2[lab == -1, 0], P2[lab == -1, 1], s=8, c="lightgrey", label="noise")
            sc = ax.scatter(P2[lab != -1, 0], P2[lab != -1, 1], s=8, c=lab[lab != -1], cmap="tab10", alpha=0.7)
            ax.set_title(f"{name} – {params[name]}")
        else:
            ax.set_title(f"{name} – no valid result"); ax.axis("on")
        ax.set_xlabel("PC1"); ax.set_ylabel("PC2")
    fig.suptitle("Clusters visualised on the first two principal components", y=1.0)
    save_fig("B9_cluster_scatter")

# %%
if RUN_RETAIL:
    # ---- Profile the clusters of the best algorithm ----
    lab = labels_all[best_overall]
    tx_sample = [tx_kept[i] for i in sel]
    overall_share = Xs.mean(axis=0)
    prof = []
    for c in sorted(set(lab)):
        m = lab == c
        share = Xs[m].mean(axis=0)
        ok = np.where(share >= 0.05)[0]
        if len(ok) == 0:
            ok = np.argsort(-share)[:5]
        ranked = ok[np.argsort(-(share[ok] / overall_share[ok]))][:5]
        prof.append({"Cluster": "noise" if c == -1 else int(c), "Transactions": int(m.sum()),
                     "Share (%)": round(m.mean() * 100, 1), "Avg basket size": round(Xs[m].sum(axis=1).mean(), 2),
                     "Distinguishing products (share in cluster)": "; ".join(f"{keep[j]} ({share[j] * 100:.0f}%)" for j in ranked)})
    profile = pd.DataFrame(prof).set_index("Cluster")
    show(profile, f"Cluster profile – {best_overall}")
    save_table(profile, "B_cluster_profile")
    R["profile"] = profile

# %% [markdown]
# ## Build the combined Word report
# The text below is generated from the results above (all numbers come from the code). **Read it, edit the wording in Word, and add your own interpretation** before you submit.

# %%
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _fmt(v):
    if isinstance(v, (float, np.floating)):
        return "–" if np.isnan(v) else f"{v:.3f}"
    return str(v)


def add_table(doc, df, index=True, font=8.5):
    d = df.reset_index() if index else df.copy()
    t = doc.add_table(rows=1, cols=len(d.columns)); t.style = "Light Grid Accent 1"
    for j, c in enumerate(d.columns):
        cell = t.rows[0].cells[j]; cell.text = ""
        r = cell.paragraphs[0].add_run(str(c)); r.bold = True; r.font.size = Pt(font)
    for _, row in d.iterrows():
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""
            cells[j].paragraphs[0].add_run(_fmt(v)).font.size = Pt(font)
    doc.add_paragraph()


def add_fig(doc, key, caption, width=6.0):
    if key in FIGS:
        doc.add_picture(FIGS[key], width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph(); r = p.add_run(caption); r.italic = True; r.font.size = Pt(9)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def para(doc, text, bold=False):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold
    return p


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def build_report(path=os.path.join(OUT_DIR, "Combined_Report.docx")):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(0.9)

    h = doc.add_heading(f"{COURSE} – Combined Assignment Report", 0)
    para(doc, f"Name: {STUDENT_NAME}    ID: {STUDENT_ID}")
    para(doc, "Part A: Comparative Analysis of Classification Algorithms for Bank Marketing Prediction\n"
              "Part B: Transaction Analysis and Clustering Algorithm Comparison for an Online Retail Store")

    # ============================ PART A ============================
    if RUN_BANK:
        res, cmt = R["results"], R["cm_table"]; best = R["best_model"]
        sc = R["results_scored"]
        doc.add_heading("PART A – Bank Marketing Classification", 1)
        doc.add_heading("A1. Business problem", 2)
        para(doc, "A Portuguese bank runs telephone campaigns to sell term deposits. The aim is to predict, before a customer is called, "
                  "whether the customer will subscribe, and to identify the most suitable classification model. Four algorithms – Logistic Regression, "
                  "Decision Tree, K-Nearest Neighbours and Random Forest – are trained and compared on the same data split.")

        doc.add_heading("A2. Data understanding and preprocessing", 2)
        rs, rc = R["bank_shape_raw"], R["bank_shape_clean"]
        para(doc, f"The dataset (semicolon-delimited CSV) contains {rs[0]:,} records and {rs[1]} columns. There are {len(R['num_cols_all'])} numerical attributes "
                  f"({', '.join(R['num_cols_all'])}), {len(R['cat_cols_all'])} categorical attributes ({', '.join(R['cat_cols_all'])}) and one target attribute (y: subscribed yes/no).")
        unk = R["unknown"]
        para(doc, f"Data quality: {R['missing_total']} true missing values were found. "
                  + (f"However several categorical columns contain the label 'unknown' ({', '.join(f'{k}: {v:,}' for k, v in unk.items())}); "
                     "this label was kept as a separate category because it may carry information. " if unk else "")
                  + f"{R['n_dups']} exact duplicate record(s) were removed, leaving {rc[0]:,} rows.")
        para(doc, f"Target distribution: {R['n_yes']:,} customers subscribed ({R['pct_yes']:.2f}%) and {R['n_no']:,} did not ({100 - R['pct_yes']:.2f}%), i.e. the classes are imbalanced.")
        para(doc, "Preprocessing: the attribute 'duration' was excluded from every model because it is only known after the call has finished and would leak the outcome. "
                  "Categorical variables were one-hot encoded and numerical variables standardised (StandardScaler). The encoder and scaler were fitted on the training data only. "
                  f"The data were split 80:20 with stratification (random_state={SEED}): {R['n_train']:,} training and {R['n_test']:,} test records, "
                  f"{R['n_features_encoded']} predictors after encoding.")

        doc.add_heading("A3. Exploratory data analysis", 2)
        add_table(doc, R["eda_counts"])
        add_fig(doc, "A1_target_distribution", "Figure A1: Distribution of the target variable.", 4.3)
        para(doc, f"Figure A1 shows that the 'no' class dominates ({100 - R['pct_yes']:.1f}%), so accuracy alone is a poor performance measure.")
        add_fig(doc, "A2_subscription_rate_by_category", "Figure A2: Subscription rate by selected customer and campaign attributes.", 6.3)
        add_fig(doc, "A3_numeric_vs_outcome", "Figure A3: Distribution of numerical attributes for subscribers and non-subscribers.", 6.3)
        add_fig(doc, "A4_subscription_by_age_group", "Figure A4: Subscription rate by age group.", 4.3)
        para(doc, "Observations from the analysis:", bold=True)
        bullets(doc, R["observations"])

        doc.add_heading("A4. Models and performance comparison", 2)
        para(doc, "The following models were trained on the same training set and evaluated on the same test set: Logistic Regression (max_iter=1000), "
                  "Decision Tree (max_depth=6, min_samples_leaf=50), KNN (k=15) and Random Forest (200 trees, min_samples_leaf=5). "
                  "Metrics treat 'subscribed' as the positive class.")
        add_table(doc, res[["Accuracy", "Precision", "Recall", "F1-score", "ROC-AUC"]].round(4))
        para(doc, "Table A1: Final performance comparison table (test set, threshold 0.5).").runs[0].italic = True
        add_fig(doc, "A5_model_comparison_grouped_bar", "Figure A5: Grouped bar chart of accuracy, precision, recall and F1-score.", 6.0)
        hi_acc, hi_rec, hi_pre = res["Accuracy"].idxmax(), res["Recall"].idxmax(), res["Precision"].idxmax()
        para(doc, f"{hi_acc} has the highest accuracy ({res.loc[hi_acc, 'Accuracy']:.3f}), but accuracy is dominated by the majority 'no' class "
                  f"(a trivial 'always no' model scores about {100 - R['pct_yes']:.1f}%). {hi_rec} has the highest recall ({res.loc[hi_rec, 'Recall']:.3f}) and "
                  f"{hi_pre} the highest precision ({res.loc[hi_pre, 'Precision']:.3f}).")

        add_fig(doc, "A6_confusion_matrices", "Figure A6: Confusion matrix of each classifier.", 5.6)
        add_table(doc, cmt)
        lowest_fp, lowest_fn = cmt["FP"].idxmin(), cmt["FN"].idxmin()
        para(doc, f"Interpretation of errors: false negatives are subscribers the bank would fail to call; false positives are calls wasted on customers who decline. "
                  f"{lowest_fn} misses the fewest subscribers (FN = {cmt.loc[lowest_fn, 'FN']}), whereas {lowest_fp} produces the fewest wasted calls (FP = {cmt.loc[lowest_fp, 'FP']}). "
                  f"Across all models the number of missed subscribers ranges from {cmt['FN'].min()} to {cmt['FN'].max()} out of {int(cmt['Actual subscribers'].iloc[0])} actual subscribers in the test set, "
                  "showing how hard it is to detect the minority class at the default 0.5 threshold.")

        add_fig(doc, "A7_roc_curves", "Figure A7: ROC curves with AUC values.", 4.8)
        hi_auc = res["ROC-AUC"].idxmax()
        para(doc, f"{hi_auc} has the best ROC-AUC ({res.loc[hi_auc, 'ROC-AUC']:.3f}), meaning it ranks subscribers above non-subscribers most reliably regardless of the threshold. "
                  "AUC values: " + ", ".join(f"{n} {v:.3f}" for n, v in res["ROC-AUC"].items()) + ".")
        add_fig(doc, "A8_cumulative_gains", "Figure A8: Cumulative gains – share of subscribers found when calling the best-scored customers first.", 4.8)

        doc.add_heading("A5. Best model and justification", 2)
        para(doc, "Selection criterion: because of the class imbalance, accuracy was not used. A missed subscriber is a lost deposit whereas a false positive costs only one phone call, "
                  "so models were ranked by the mean of ROC-AUC, F1-score and Recall.")
        add_table(doc, sc.sort_values("Selection score", ascending=False)[["Precision", "Recall", "F1-score", "ROC-AUC", "Selection score"]].round(4))
        para(doc, f"Best model: {best}. It obtains ROC-AUC {sc.loc[best, 'ROC-AUC']:.3f}, F1 {sc.loc[best, 'F1-score']:.3f}, recall {sc.loc[best, 'Recall']:.3f} and "
                  f"precision {sc.loc[best, 'Precision']:.3f}, giving the highest selection score ({sc.loc[best, 'Selection score']:.3f}). "
                  f"It finds {cmt.loc[best, 'Subscribers found (recall %)']:.1f}% of the real subscribers at the default threshold, "
                  f"with {cmt.loc[best, 'Wasted calls (FP)']:,} wasted calls.")
        rb = R["results_bal"]
        para(doc, "Extension – class_weight='balanced': giving more weight to the minority class raises recall at the cost of precision (Table A2), which shows the trade-off available to the bank.")
        add_table(doc, rb.round(4))
        para(doc, "Table A2: Performance with class-weight balancing (KNN unchanged).").runs[0].italic = True

        doc.add_heading("A6. Business recommendations and limitations", 2)
        para(doc, "Business recommendations", bold=True)
        bullets(doc, [
            f"Rank customers by predicted subscription probability before each campaign and call the highest-scoring customers first. With {best}, contacting the top 20% of customers "
            f"would capture {R['capture20'] * 100:.1f}% of all subscribers (a {R['lift20']:.2f}x lift over random calling), reducing call-centre cost per deposit.",
            f"Choose the decision threshold from business costs rather than using 0.5. A threshold of {R['thr60']:.2f} would find about {R['rec60'] * 100:.0f}% of subscribers "
            f"with {R['prec60'] * 100:.0f}% precision while contacting {R['frac60'] * 100:.0f}% of customers; the bank can raise or lower it depending on call cost and expected deposit value.",
            "Use the segments with the highest subscription rates found in the EDA (Section A3) to design targeted offers, and retrain / monitor the model for every new campaign."])
        para(doc, "Limitations", bold=True)
        bullets(doc, [
            "Class imbalance and a fixed 0.5 threshold: with only about "
            f"{R['pct_yes']:.0f}% positives the models are biased towards 'no', so reported recall is low; without cost-sensitive learning or resampling, many subscribers remain undetected and accuracy overstates quality.",
            "Single train/test split and limited hyper-parameter tuning: results can change with a different split or tuned parameters (especially for KNN, the Decision Tree and the Random Forest), so differences between close models may not be statistically significant. Cross-validation would give more reliable estimates.",
            "Historical data from one bank and period: customer behaviour and economic conditions change, and excluding 'duration' (necessary to avoid leakage) lowers predictive power; the model may degrade on future campaigns."])

    # ============================ PART B ============================
    if RUN_RETAIL:
        doc.add_page_break()
        info = R["retail_info"]; comp = R["comp"]
        doc.add_heading("PART B – Online Retail Transaction Clustering", 1)
        doc.add_heading("B1. Problem and data", 2)
        para(doc, "The online retail store holds transaction records in which several products are bought together, but has no systematic view of what its baskets look like "
                  "or whether groups of similar baskets exist. The objective is to describe the transaction data, represent every transaction numerically, reduce dimensionality with PCA, "
                  "and compare K-means, hierarchical, DBSCAN and spectral clustering with the silhouette coefficient and the Davies–Bouldin index.")
        para(doc, f"The file '{info['file']}' ({info['format']} format) contained {info['raw_rows']:,} raw rows. After cleaning with Pandas/Python (trimming and lower-casing product names, "
                  f"removing empty cells and duplicate items within a basket, and dropping empty transactions"
                  + (", cancelled invoices and non-product lines" if info['format'] == 'long' else "")
                  + f") there are {R['n_tx']:,} transactions and {R['n_items']:,} distinct products.")

        doc.add_heading("B2. Descriptive statistics and visualisation", 2)
        bs = R["basket_stats"]
        para(doc, f"Transaction size: mean {bs['mean']:.2f} items, median {bs['50%']:.0f}, standard deviation {bs['std']:.2f}, minimum {bs['min']:.0f}, maximum {bs['max']:.0f}; "
                  f"the most common basket size is {R['mode_size']} item(s) and {R['single_item_share']:.1f}% of transactions contain a single item.")
        add_table(doc, R["top_items"], index=False)
        para(doc, "Table B1: The 15 most frequently purchased products.").runs[0].italic = True
        add_fig(doc, "B1_top_products", "Figure B1: Top 15 most frequently purchased products.", 5.2)
        add_fig(doc, "B2_transaction_size_distribution", "Figure B2: Distribution of transaction size.", 4.8)
        t0 = R["top_items"].iloc[0]
        para(doc, f"Interpretation: '{t0['Product']}' is the most frequent product (in {t0['Support (%)']:.1f}% of transactions). "
                  + ("The transaction-size histogram is right-skewed: most customers buy a few items while a small number buy large baskets." if bs['mean'] > bs['50%'] else
                     "The transaction-size histogram shows that most baskets are of similar, small size."))

        doc.add_heading("B3. Feature representation and PCA", 2)
        para(doc, f"Each transaction was represented as a binary vector over the {R['feat_shape'][1]} most frequent products (these cover {R['coverage']:.1f}% of all item occurrences); "
                  f"a random sample of {R['sample_shape'][0]:,} transactions was clustered because hierarchical and spectral clustering require pairwise distances. "
                  f"PCA shows that the first two components explain only {R['var2']:.1f}% of the variance (PC1 {R['pc1']:.1f}%, PC2 {R['pc2']:.1f}%); {R['n_pca']} components retain {R['varC']:.1f}% "
                  f"and are used as the clustering space, while {R['n_for_80']} components would be needed for 80%. This indicates that baskets are sparse and spread over many dimensions.")
        add_fig(doc, "B3_pca_variance", "Figure B3: PCA scree plot and cumulative explained variance.", 6.2)
        add_fig(doc, "B4_pca_scatter", "Figure B4: Transactions on the first two principal components.", 4.3)

        doc.add_heading("B4. Clustering results", 2)
        para(doc, f"K-means was run for k = 2 to 8; the best silhouette was obtained for k = {R['k_best']}, which was then also used for hierarchical (Ward) and spectral clustering. "
                  "DBSCAN was tuned by a grid search over eps and min_samples (best silhouette with at least two clusters and at most 50% noise).")
        add_fig(doc, "B5_kmeans_selection", "Figure B5: K-means elbow, silhouette and Davies–Bouldin versus k.", 6.3)
        add_fig(doc, "B6_dendrogram", "Figure B6: Truncated Ward dendrogram.", 5.8)
        add_fig(doc, "B7_dbscan_kdistance", "Figure B7: DBSCAN k-distance plot.", 4.3)
        add_table(doc, comp[["Parameters", "Clusters", "Noise points", "Silhouette (higher=better)", "Davies-Bouldin (lower=better)"]].round(4))
        para(doc, "Table B2: Clustering comparison (metrics computed in the 10-component PCA space; DBSCAN noise excluded).").runs[0].italic = True
        add_fig(doc, "B8_clustering_metrics", "Figure B8: Silhouette and Davies–Bouldin scores per algorithm.", 5.8)
        add_fig(doc, "B9_cluster_scatter", "Figure B9: Cluster assignments shown on the first two principal components.", 6.0)

        doc.add_heading("B5. Discussion", 2)
        bs_sil = comp.loc[R["best_sil"], "Silhouette (higher=better)"]
        strength = "strong" if bs_sil > 0.5 else "reasonable" if bs_sil > 0.25 else "weak"
        para(doc, f"{R['best_sil']} achieved the highest silhouette coefficient ({bs_sil:.3f}) and {R['best_dbi']} the lowest Davies–Bouldin index "
                  f"({comp.loc[R['best_dbi'], 'Davies-Bouldin (lower=better)']:.3f}); using the average rank of both measures, {R['best_overall']} is the best algorithm for this data. "
                  f"A silhouette of {bs_sil:.2f} indicates {strength} cluster structure"
                  + (" – the baskets form overlapping groups rather than clearly separated clusters." if strength == "weak" else ".")
                  + " Note that DBSCAN's scores exclude noise points, so they are not fully comparable with the other algorithms.")
        para(doc, "Profile of the clusters found by the best algorithm:")
        add_table(doc, R["profile"], font=8)

        doc.add_heading("B6. Limitations and business implications", 2)
        para(doc, "Limitations", bold=True)
        bullets(doc, [
            "Binary representation ignores quantities, prices and time; only the top products were kept and a sample of transactions was clustered, so rare products and some baskets are not represented.",
            f"PCA with {R['n_pca']} components keeps only {R['varC']:.1f}% of the variance, and sparse binary data have no natural spherical clusters; internal measures (silhouette, Davies–Bouldin) judge geometric separation, not whether clusters are useful for the business.",
            "Results depend on parameters (k, eps, min_samples, number of neighbours) and the random seed; DBSCAN in particular is sensitive to eps in sparse high-dimensional data."])
        para(doc, "Business implications", bold=True)
        bullets(doc, [
            "Frequently bought products and distinguishing products of each basket group can be used for product bundling, shelf/website placement and cross-selling recommendations.",
            "Basket groups can support targeted promotions and stock planning (e.g. keeping stock of products that dominate the largest groups); the groups should be validated with association rules or A/B tests before large investments."])

    doc.add_heading("Conclusion", 1)
    concl = []
    if RUN_BANK:
        concl.append(f"For bank marketing, {R['best_model']} was selected as the best model using ROC-AUC, F1 and recall rather than accuracy, and can be used to rank customers for telephone campaigns.")
    if RUN_RETAIL:
        concl.append(f"For the retail data, {R['best_overall']} gave the best overall clustering quality, although the cluster structure is {('weak' if comp['Silhouette (higher=better)'].max() <= 0.25 else 'moderate')} and should be interpreted with care.")
    para(doc, " ".join(concl))
    para(doc, "Code: see bank_retail_analysis.py / bank_retail_analysis.ipynb in the accompanying GitHub repository.")
    doc.save(path)
    return path


report_path = build_report()
print("Report saved to:", report_path)
print("Figures:", len(FIGS), "| Tables folder:", TAB_DIR)